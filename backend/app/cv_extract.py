"""Turn an uploaded CV file into plain text.

Deliberately a small, separate implementation rather than an import from
Sabha's council/extract.py, even though that module does this well: the
two repos are independently cloneable and coupling them through a
filesystem path would break both. It uses the same two libraries so the
text comes out the same, and keeps the one guard that really matters —
an image-only PDF extracts to nothing, and saying so precisely is far
better than storing an empty CV that silently screens every lead as a
poor match.

Bureau doesn't reproduce Sabha's structural ATS signals (tables, text
boxes, embedded images). Those exist to audit how a parser will mangle a
CV, which is Sabha's job; Bureau only needs the words.

Extraction happens in memory. The uploaded bytes are never written to
disk — only the extracted text is, to data/cv.txt.
"""

import io
import re

MAX_UPLOAD_BYTES = 10 * 1024 * 1024
MIN_TEXT_CHARS = 120


class ExtractionError(Exception):
    """Raised with a message intended to be shown to the person directly."""


def _clean(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = text.replace("\xa0", " ").replace("•", "- ").replace("●", "- ")
    text = re.sub(r"[ \t]+\n", "\n", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _from_pdf(data: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:
        raise ExtractionError("PDF support needs pypdf installed on the server.") from exc
    try:
        reader = PdfReader(io.BytesIO(data))
        text = _clean("\n\n".join((page.extract_text() or "") for page in reader.pages))
    except Exception as exc:
        raise ExtractionError(f"Could not read that PDF ({type(exc).__name__}).") from exc

    if len(text) < MIN_TEXT_CHARS:
        raise ExtractionError(
            "That PDF has almost no extractable text — it looks like a scan or an "
            "image-only export. Applicant tracking systems would read it the same "
            "way, as blank. Export a text-based PDF, or paste the text instead."
        )
    return text


def _from_docx(data: bytes) -> str:
    try:
        import docx
    except ImportError as exc:
        raise ExtractionError("Word support needs python-docx installed on the server.") from exc
    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise ExtractionError(f"Could not read that Word file ({type(exc).__name__}).") from exc

    parts = [p.text for p in document.paragraphs]
    # CVs laid out in tables are common, and their content is invisible if
    # only paragraphs are read.
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append("\t".join(cells))

    text = _clean("\n".join(parts))
    if len(text) < MIN_TEXT_CHARS:
        raise ExtractionError("That Word file has almost no readable text in it.")
    return text


def from_upload(filename: str, data: bytes) -> tuple[str, str]:
    """Returns (text, detected_kind). Sniffs magic bytes as well as the
    extension, since a CV emailed around for years often arrives with the
    wrong one."""
    if not data:
        raise ExtractionError("That file is empty.")
    if len(data) > MAX_UPLOAD_BYTES:
        raise ExtractionError("That file is larger than 10MB — a CV shouldn't be.")

    name = (filename or "").lower()
    if name.endswith(".pdf") or data[:5] == b"%PDF-":
        return _from_pdf(data), "pdf"
    if name.endswith(".docx") or data[:2] == b"PK":
        return _from_docx(data), "docx"
    if name.endswith(".doc"):
        raise ExtractionError(
            "Old-style .doc files can't be read here. Save it as .docx or PDF and try again."
        )

    try:
        text = _clean(data.decode("utf-8", errors="strict"))
    except UnicodeDecodeError as exc:
        raise ExtractionError(
            "That doesn't look like a PDF, a Word file, or plain text."
        ) from exc
    if len(text) < MIN_TEXT_CHARS:
        raise ExtractionError("That file is too short to be a CV.")
    return text, "text"
